import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def create_word_document(content):
    doc = docx.Document()
    
    # Set up styles
    styles = doc.styles
    title_style = styles.add_style('CustomTitle', docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    title_style.font.size = Pt(16)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 112, 192)  # Blue color
    
    heading_style = styles.add_style('CustomHeading', docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    heading_style.font.size = Pt(14)
    heading_style.font.bold = True
    
    subheading_style = styles.add_style('CustomSubheading', docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    subheading_style.font.size = Pt(11)
    subheading_style.font.bold = True
    
    normal_style = styles.add_style('CustomNormal', docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    normal_style.font.size = Pt(11)
    
    # Add title
    title = doc.add_paragraph("Sanction Team API Documentation", style='CustomTitle')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Process content
    lines = content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        if line.startswith(('1.', '2.', '3.', '4.', '5.')):
            p = doc.add_paragraph(line, style='CustomHeading')
            i += 1
            
            while i < len(lines) and lines[i].strip() != '':
                p = doc.add_paragraph(lines[i].strip(), style='CustomNormal')
                i += 1
        
        elif line in ('Endpoint:', 'Method:', 'Middleware:', 'Description:', 'Query Parameters:', 'Body Parameters:', 'Response:', 'Error Handling'):
            p = doc.add_paragraph(line, style='CustomSubheading')
            i += 1
            
            while i < len(lines) and lines[i].strip() != '':
                if lines[i].strip().startswith(('200 OK', '400 Bad Request', '403 Forbidden', '404 Not Found', '500 Internal Server Error')):
                    p = doc.add_paragraph(lines[i].strip(), style='CustomSubheading')
                else:
                    p = doc.add_paragraph(lines[i].strip(), style='CustomNormal')
                p.paragraph_format.left_indent = Pt(36)
                i += 1
        
        elif line == 'json':
            p = doc.add_paragraph()
            run = p.add_run(lines[i+1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            p.paragraph_format.left_indent = Pt(36)
            i += 2
            
            while i < len(lines) and lines[i].strip() != '':
                p = doc.add_paragraph(lines[i].strip())
                p.style = doc.styles['Normal']
                p.style.font.name = 'Courier New'
                p.style.font.size = Pt(9)
                p.paragraph_format.left_indent = Pt(36)
                i += 1
        
        else:
            p = doc.add_paragraph(line, style='CustomNormal')
            i += 1
    
    # Save the document
    doc.save('Sanction_Team_API_Documentation.docx')

# Content from the provided documentation
content = """1. Get Sanction Team

Retrieves a list of all members of the "Sanction Team," along with a count of assigned and converted applications.

    Endpoint: /sanction-team
    Method: GET
    Middleware: verifySanctionsHead
    Description: Fetches the sanction team members with pagination and application statistics (assigned and converted applications).

Query Parameters:

    pageNumber (optional, number, default: 1): Specifies the page of results.
    pageSize (optional, number, default: 50): Specifies the number of results per page.

Response:

    200 OK

    json

{
  "message": "Here are your team members",
  "teamWithCounts": [
    {
      "_id": "adminId",
      "name": "Admin Name",
      "applicationsAssignedCount": 10,
      "applicationsConvertedCount": 5
    },
    ...
  ]
}

500 Internal Server Error

json

    {
      "message": "Internal Error Occured"
    }

2. Get Sanction Team Member Details

Retrieves details of a specific sanction team member along with their deviation count.

    Endpoint: /get-team-member
    Method: POST
    Middleware: verifySanctionsHeadWithTeam
    Description: Fetches a specific sanction team member's information and the number of deviations they created.

Body Parameters:

    adminId (required, string): The ID of the sanction team member.

Response:

    200 OK

    json

{
  "message": "Here is the team member for you",
  "sanctionTeamMember": {
    "_id": "adminId",
    "name": "Admin Name",
    "adminRole": "SanctionT",
    ...
  },
  "deviationCount": 5
}

403 Forbidden

json

{
  "message": "You are not authorized to view this resource"
}

404 Not Found

json

{
  "message": "Admin not found"
}

500 Internal Server Error

json

    {
      "message": "Internal servre error"
    }

3. Get Assigned Applications Pending by Status

Retrieves pending applications assigned to a specific sanction team member based on their status.

    Endpoint: /get-member-application
    Method: POST
    Middleware: verifySanctionsHead
    Description: Fetches the applications assigned to a sanction team member with the specified status (e.g., pending, approved).

Query Parameters:

    status (required, string): Comma-separated values of loan statuses (e.g., "Pending_Approved").
    pageNumber (optional, number, default: 1): Specifies the page of results.
    pageSize (optional, number, default: 50): Specifies the number of results per page.

Body Parameters:

    adminId (required, string): The ID of the sanction team member.

Response:

    200 OK

    json

{
  "message": "Here are your assigned loans",
  "assignedLoans": [
    {
      "_id": "loanId",
      "status": "Pending",
      "assignedTo": "adminId",
      ...
    },
    ...
  ]
}

400 Bad Request

json

{
  "message": "Status param is required"
}

404 Not Found

json

{
  "message": "Admin not found"
}

500 Internal Server Error

json

    {
      "message": "Internal Server Error"
    }

4. Assign Application to Admin

Assigns a loan application to a specific sanction team member.

    Endpoint: /assign-applicaiton
    Method: POST
    Middleware: verifySanctionsHead
    Description: Assigns a loan application to a sanction team member by updating both the member's record and the loan application.

Body Parameters:

    adminId (required, string): The ID of the admin to whom the loan is assigned.
    applicationId (required, string): The ID of the loan application being assigned.

Response:

    200 OK

    json

{
  "message": "Application assigned successfully",
  "admin": {
    "_id": "adminId",
    "applicationsassigned": ["applicationId"]
  },
  "loan": {
    "_id": "applicationId",
    "assignedTo": "adminId"
  }
}

400 Bad Request

json

{
  "message": "AdminId and ApplicationId are required"
}

404 Not Found

json

{
  "message": "Admin or Loan not found"
}

500 Internal Server Error

json

    {
      "message": "Internal server error"
    }

5. Initiate Third-Party Sanction Process

Initiates a third-party sanction contract via the Signzy service.

    Endpoint: /initiate-sanction
    Method: POST
    Middleware: verifyToken
    Description: Initiates a contract with a third-party service to process sanctions.

Response:

    200 OK

    json

{
  "message": "Third-party sanction process initiated successfully"
}

500 Internal Server Error

json

    {
      "message": "Internal server error"
    }

Error Handling

In case of errors, the API will return appropriate error messages with HTTP status codes. Common error status codes include:

    400 Bad Request: For invalid input parameters.
    403 Forbidden: When a user is not authorized to access a specific resource.
    404 Not Found: When the requested resource is not found.
    500 Internal Server Error: When there is a problem with the server."""

create_word_document(content)

print("Word document 'Sanction_Team_API_Documentation.docx' has been created.")